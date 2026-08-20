"""将比赛 Excel 数据导入 SQLite，并剔除所有包含图片的完整会话。"""

from __future__ import annotations

import argparse
import os
import sqlite3
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = PROJECT_ROOT / "赛题 1：数据共情者-业务数据.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "data" / "customer_service.db"


TABLE_SCHEMAS = {
    "chat_messages": """
        CREATE TABLE chat_messages (
            message_id TEXT PRIMARY KEY,
            session_id TEXT NOT NULL,
            message_seq INTEGER NOT NULL,
            sent_at TEXT,
            role TEXT,
            buyer_nickname TEXT,
            sender TEXT,
            store TEXT,
            scene_major TEXT,
            scene_minor TEXT,
            is_target_buyer_message INTEGER NOT NULL DEFAULT 0,
            message_text TEXT,
            content_type TEXT,
            chat_content TEXT,
            category TEXT,
            image_path TEXT,
            related_order_id TEXT,
            related_ticket_id TEXT
        )
    """,
    "orders": """
        CREATE TABLE orders (
            order_id TEXT PRIMARY KEY,
            session_id TEXT,
            buyer_nickname TEXT,
            store TEXT,
            sku TEXT,
            product_name TEXT,
            quantity INTEGER,
            unit_price REAL,
            paid_amount REAL,
            order_status TEXT,
            ordered_at TEXT,
            paid_at TEXT,
            shipped_at TEXT,
            carrier TEXT,
            tracking_no TEXT,
            province TEXT,
            city TEXT,
            gift TEXT,
            buyer_note TEXT
        )
    """,
    "reship_exchange_tickets": """
        CREATE TABLE reship_exchange_tickets (
            ticket_id TEXT PRIMARY KEY,
            session_id TEXT,
            related_order_id TEXT,
            buyer_nickname TEXT,
            store TEXT,
            ticket_type TEXT,
            aftersales_reason TEXT,
            outgoing_sku TEXT,
            outgoing_product_name TEXT,
            quantity INTEGER,
            original_tracking_no TEXT,
            replacement_tracking_no TEXT,
            carrier TEXT,
            warehouse TEXT,
            urgent TEXT,
            ticket_status TEXT,
            assignee TEXT,
            created_at TEXT,
            completed_at TEXT
        )
    """,
    "offline_payment_tickets": """
        CREATE TABLE offline_payment_tickets (
            ticket_id TEXT PRIMARY KEY,
            session_id TEXT,
            related_order_id TEXT,
            buyer_nickname TEXT,
            store TEXT,
            payment_type TEXT,
            refund_issue_type TEXT,
            refund_amount REAL,
            account_real_name TEXT,
            payment_account TEXT,
            related_tracking_no TEXT,
            transfer_status TEXT,
            ticket_status TEXT,
            assignee TEXT,
            created_at TEXT,
            completed_at TEXT
        )
    """,
    "logistics_tickets": """
        CREATE TABLE logistics_tickets (
            ticket_id TEXT PRIMARY KEY,
            session_id TEXT,
            related_order_id TEXT,
            buyer_nickname TEXT,
            store TEXT,
            issue_type TEXT,
            carrier TEXT,
            tracking_no TEXT,
            warehouse TEXT,
            paid_amount REAL,
            resolution TEXT,
            province TEXT,
            city TEXT,
            ticket_status TEXT,
            assignee TEXT,
            created_at TEXT,
            completed_at TEXT
        )
    """,
    "adverse_reaction_tickets": """
        CREATE TABLE adverse_reaction_tickets (
            ticket_id TEXT PRIMARY KEY,
            session_id TEXT,
            related_order_id TEXT,
            buyer_nickname TEXT,
            store TEXT,
            channel_type TEXT,
            age INTEGER,
            skin_type TEXT,
            product_name TEXT,
            batch_no TEXT,
            body_area TEXT,
            symptom_description TEXT,
            onset_after TEXT,
            stopped_use TEXT,
            sought_medical_attention TEXT,
            task_status TEXT,
            assignee TEXT,
            created_at TEXT,
            completed_at TEXT
        )
    """,
    "aftersales_return_tickets": """
        CREATE TABLE aftersales_return_tickets (
            ticket_id TEXT PRIMARY KEY,
            session_id TEXT,
            related_order_id TEXT,
            buyer_nickname TEXT,
            store TEXT,
            parcel_type TEXT,
            return_reason TEXT,
            return_tracking_no TEXT,
            carrier TEXT,
            refund_id TEXT,
            receipt_advice TEXT,
            is_abnormal TEXT,
            task_status TEXT,
            assignee TEXT,
            created_at TEXT,
            completed_at TEXT
        )
    """,
}


SHEET_MAPPINGS = {
    "订单": (
        "orders",
        {
            "订单号": "order_id",
            "会话ID": "session_id",
            "买家昵称": "buyer_nickname",
            "店铺": "store",
            "商品货号": "sku",
            "商品名称": "product_name",
            "数量": "quantity",
            "单价(元)": "unit_price",
            "实付金额(元)": "paid_amount",
            "订单状态": "order_status",
            "下单时间": "ordered_at",
            "付款时间": "paid_at",
            "发货时间": "shipped_at",
            "快递公司": "carrier",
            "物流单号": "tracking_no",
            "收货省": "province",
            "收货市": "city",
            "赠品": "gift",
            "买家留言": "buyer_note",
        },
    ),
    "补发换货工单": (
        "reship_exchange_tickets",
        {
            "工单号": "ticket_id",
            "会话ID": "session_id",
            "关联订单号": "related_order_id",
            "买家昵称": "buyer_nickname",
            "店铺": "store",
            "工单类型": "ticket_type",
            "售后原因": "aftersales_reason",
            "发出商品货号": "outgoing_sku",
            "发出商品名称": "outgoing_product_name",
            "数量": "quantity",
            "原订单物流单号": "original_tracking_no",
            "补发物流单号": "replacement_tracking_no",
            "快递公司": "carrier",
            "发货仓库": "warehouse",
            "客诉加急": "urgent",
            "工单状态": "ticket_status",
            "处理人": "assignee",
            "创建时间": "created_at",
            "完成时间": "completed_at",
        },
    ),
    "线下打款工单": (
        "offline_payment_tickets",
        {
            "工单号": "ticket_id",
            "会话ID": "session_id",
            "关联订单号": "related_order_id",
            "买家昵称": "buyer_nickname",
            "店铺": "store",
            "打款类型": "payment_type",
            "退款问题类型": "refund_issue_type",
            "退款金额(元)": "refund_amount",
            "支付宝实名": "account_real_name",
            "支付宝账号": "payment_account",
            "相关物流单号": "related_tracking_no",
            "转账状态": "transfer_status",
            "工单状态": "ticket_status",
            "处理人": "assignee",
            "创建时间": "created_at",
            "完成时间": "completed_at",
        },
    ),
    "物流工单": (
        "logistics_tickets",
        {
            "工单号": "ticket_id",
            "会话ID": "session_id",
            "关联订单号": "related_order_id",
            "买家昵称": "buyer_nickname",
            "店铺": "store",
            "问题类型": "issue_type",
            "快递公司": "carrier",
            "问题包裹物流单号": "tracking_no",
            "发货仓": "warehouse",
            "订单实付(元)": "paid_amount",
            "处理方案": "resolution",
            "收货省": "province",
            "收货市": "city",
            "工单状态": "ticket_status",
            "处理人": "assignee",
            "创建时间": "created_at",
            "完成时间": "completed_at",
        },
    ),
    "不良反应工单": (
        "adverse_reaction_tickets",
        {
            "工单号": "ticket_id",
            "会话ID": "session_id",
            "关联订单号": "related_order_id",
            "买家昵称": "buyer_nickname",
            "店铺": "store",
            "类型": "channel_type",
            "年龄": "age",
            "肤质": "skin_type",
            "使用商品": "product_name",
            "产品批次号": "batch_no",
            "不适部位": "body_area",
            "症状描述": "symptom_description",
            "用后多久出现": "onset_after",
            "是否停用": "stopped_use",
            "是否就医": "sought_medical_attention",
            "任务状态": "task_status",
            "处理人": "assignee",
            "创建时间": "created_at",
            "完成时间": "completed_at",
        },
    ),
    "售后退货工单": (
        "aftersales_return_tickets",
        {
            "工单号": "ticket_id",
            "会话ID": "session_id",
            "关联订单号": "related_order_id",
            "买家昵称": "buyer_nickname",
            "店铺": "store",
            "包裹类型": "parcel_type",
            "退货原因": "return_reason",
            "退货物流单号": "return_tracking_no",
            "快递公司": "carrier",
            "退款编号": "refund_id",
            "签收建议": "receipt_advice",
            "是否异常": "is_abnormal",
            "任务状态": "task_status",
            "处理人": "assignee",
            "创建时间": "created_at",
            "完成时间": "completed_at",
        },
    ),
}


CHAT_MAPPING = {
    "message_id": "message_id",
    "会话ID": "session_id",
    "消息序号": "message_seq",
    "发送时间": "sent_at",
    "角色": "role",
    "买家昵称": "buyer_nickname",
    "发送方": "sender",
    "店铺": "store",
    "scene_major": "scene_major",
    "scene_minor": "scene_minor",
    "is_target_buyer_message": "is_target_buyer_message",
    "message_text": "message_text",
    "内容类型": "content_type",
    "chat_content": "chat_content",
    "category": "category",
    "image_path": "image_path",
    "关联订单号": "related_order_id",
    "关联工单号": "related_ticket_id",
}


def normalize_value(value: Any) -> Any:
    """将 Excel 值转换为 SQLite 可稳定保存的基础类型。"""
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if value is None:
        return None
    if isinstance(value, bool):
        return int(value)
    return value


def sheet_records(worksheet: Any) -> list[dict[str, Any]]:
    """按首行表头读取工作表，忽略完全空白的行。"""
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    records: list[dict[str, Any]] = []
    for row in rows:
        if not any(value is not None for value in row):
            continue
        records.append(
            {
                header: normalize_value(value)
                for header, value in zip(headers, row)
                if header
            }
        )
    return records


def contains_image(record: dict[str, Any]) -> bool:
    """依据内容类型或图片路径识别图片消息。"""
    content_type = str(record.get("内容类型") or "").lower()
    image_path = str(record.get("image_path") or "").strip()
    return any(marker in content_type for marker in ("图片", "照片", "image")) or bool(
        image_path
    )


def create_schema(connection: sqlite3.Connection) -> None:
    connection.execute("PRAGMA foreign_keys = ON")
    # 数据库采用离线一次性构建，DELETE 日志模式可确保关闭连接后得到单一完整文件。
    connection.execute("PRAGMA journal_mode = DELETE")

    connection.execute(
        """
        CREATE TABLE import_metadata (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL
        )
        """
    )
    connection.execute(
        """
        CREATE TABLE conversations (
            session_id TEXT PRIMARY KEY,
            buyer_nickname TEXT,
            store TEXT,
            scene_major TEXT,
            scene_minor TEXT,
            first_message_at TEXT,
            last_message_at TEXT,
            message_count INTEGER NOT NULL,
            buyer_message_count INTEGER NOT NULL,
            customer_service_message_count INTEGER NOT NULL
        )
        """
    )
    for ddl in TABLE_SCHEMAS.values():
        connection.execute(ddl)
    # 仿真回复与原始聊天记录分表保存，重新导入时可得到干净的演示环境。
    connection.execute(
        """
        CREATE TABLE simulation_messages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            session_id TEXT NOT NULL,
            role TEXT NOT NULL DEFAULT '客服' CHECK (role IN ('客服', '买家')),
            message_text TEXT NOT NULL CHECK (length(message_text) BETWEEN 1 AND 500),
            created_at TEXT NOT NULL,
            FOREIGN KEY (session_id) REFERENCES conversations(session_id)
        )
        """
    )


def insert_records(
    connection: sqlite3.Connection,
    table: str,
    records: Iterable[dict[str, Any]],
    mapping: dict[str, str],
) -> int:
    columns = list(mapping.values())
    placeholders = ", ".join("?" for _ in columns)
    sql = f"INSERT INTO {table} ({', '.join(columns)}) VALUES ({placeholders})"
    values = [
        tuple(normalize_value(record.get(source)) for source in mapping)
        for record in records
    ]
    connection.executemany(sql, values)
    return len(values)


def build_conversations(connection: sqlite3.Connection) -> None:
    """从过滤后的消息表生成可供左侧会话列表使用的聚合表。"""
    connection.execute(
        """
        INSERT INTO conversations (
            session_id,
            buyer_nickname,
            store,
            scene_major,
            scene_minor,
            first_message_at,
            last_message_at,
            message_count,
            buyer_message_count,
            customer_service_message_count
        )
        SELECT
            session_id,
            MAX(buyer_nickname),
            MAX(store),
            MAX(scene_major),
            MAX(scene_minor),
            MIN(sent_at),
            MAX(sent_at),
            COUNT(*),
            SUM(CASE WHEN role = '买家' THEN 1 ELSE 0 END),
            SUM(CASE WHEN role = '客服' THEN 1 ELSE 0 END)
        FROM chat_messages
        GROUP BY session_id
        """
    )


def create_indexes_and_views(connection: sqlite3.Connection) -> None:
    statements = [
        "CREATE INDEX idx_chat_session_seq ON chat_messages(session_id, message_seq)",
        "CREATE INDEX idx_chat_sent_at ON chat_messages(sent_at)",
        "CREATE INDEX idx_chat_order ON chat_messages(related_order_id)",
        "CREATE INDEX idx_chat_ticket ON chat_messages(related_ticket_id)",
        "CREATE INDEX idx_orders_session ON orders(session_id)",
        "CREATE INDEX idx_orders_buyer ON orders(buyer_nickname)",
        "CREATE INDEX idx_orders_tracking ON orders(tracking_no)",
        "CREATE INDEX idx_simulation_messages_session_created ON simulation_messages(session_id, created_at, id)",
    ]
    for table in (
        "reship_exchange_tickets",
        "offline_payment_tickets",
        "logistics_tickets",
        "adverse_reaction_tickets",
        "aftersales_return_tickets",
    ):
        statements.extend(
            [
                f"CREATE INDEX idx_{table}_session ON {table}(session_id)",
                f"CREATE INDEX idx_{table}_order ON {table}(related_order_id)",
                f"CREATE INDEX idx_{table}_buyer ON {table}(buyer_nickname)",
            ]
        )
    for statement in statements:
        connection.execute(statement)

    # 统一工单视图便于 Agent 按会话查询全部售后记录。
    connection.execute(
        """
        CREATE VIEW ticket_overview AS
        SELECT ticket_id, session_id, related_order_id, buyer_nickname,
               '补发换货' AS ticket_category, ticket_type AS issue_type,
               ticket_status AS status, created_at, completed_at
        FROM reship_exchange_tickets
        UNION ALL
        SELECT ticket_id, session_id, related_order_id, buyer_nickname,
               '线下打款', refund_issue_type, ticket_status, created_at, completed_at
        FROM offline_payment_tickets
        UNION ALL
        SELECT ticket_id, session_id, related_order_id, buyer_nickname,
               '物流', issue_type, ticket_status, created_at, completed_at
        FROM logistics_tickets
        UNION ALL
        SELECT ticket_id, session_id, related_order_id, buyer_nickname,
               '不良反应', symptom_description, task_status, created_at, completed_at
        FROM adverse_reaction_tickets
        UNION ALL
        SELECT ticket_id, session_id, related_order_id, buyer_nickname,
               '售后退货', return_reason, task_status, created_at, completed_at
        FROM aftersales_return_tickets
        """
    )


def write_metadata(
    connection: sqlite3.Connection,
    source: Path,
    total_chat_rows: int,
    excluded_session_count: int,
    excluded_message_count: int,
    retained_message_count: int,
) -> None:
    metadata = {
        "source_file": source.name,
        "imported_at": datetime.now().isoformat(timespec="seconds"),
        "image_filter_policy": "删除任何包含图片消息的完整会话",
        "source_chat_message_count": str(total_chat_rows),
        "excluded_session_count": str(excluded_session_count),
        "excluded_chat_message_count": str(excluded_message_count),
        "retained_chat_message_count": str(retained_message_count),
    }
    connection.executemany(
        "INSERT INTO import_metadata (key, value) VALUES (?, ?)", metadata.items()
    )


def import_workbook(source: Path, output: Path) -> dict[str, int]:
    if not source.is_file():
        raise FileNotFoundError(f"找不到源文件：{source}")

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary_output = output.with_suffix(output.suffix + ".tmp")
    if temporary_output.exists():
        temporary_output.unlink()

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        chat_records = sheet_records(workbook["聊天记录"])
        image_session_ids = {
            str(record["会话ID"])
            for record in chat_records
            if record.get("会话ID") and contains_image(record)
        }
        retained_chat_records = [
            record
            for record in chat_records
            if str(record.get("会话ID")) not in image_session_ids
        ]
        excluded_message_count = len(chat_records) - len(retained_chat_records)

        counts: Counter[str] = Counter()
        connection = sqlite3.connect(temporary_output)
        try:
            create_schema(connection)
            counts["chat_messages"] = insert_records(
                connection, "chat_messages", retained_chat_records, CHAT_MAPPING
            )
            build_conversations(connection)
            counts["conversations"] = connection.execute(
                "SELECT COUNT(*) FROM conversations"
            ).fetchone()[0]

            for sheet_name, (table_name, mapping) in SHEET_MAPPINGS.items():
                counts[table_name] = insert_records(
                    connection,
                    table_name,
                    sheet_records(workbook[sheet_name]),
                    mapping,
                )

            create_indexes_and_views(connection)
            write_metadata(
                connection,
                source,
                total_chat_rows=len(chat_records),
                excluded_session_count=len(image_session_ids),
                excluded_message_count=excluded_message_count,
                retained_message_count=len(retained_chat_records),
            )
            connection.execute("ANALYZE")
            connection.execute("PRAGMA optimize")
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            # 必须在原子替换数据库文件之前关闭连接，确保所有页均已写入主文件。
            connection.close()
    finally:
        workbook.close()

    os.replace(temporary_output, output)
    counts["excluded_sessions"] = len(image_session_ids)
    counts["excluded_messages"] = excluded_message_count
    return dict(counts)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    counts = import_workbook(args.source.resolve(), args.output.resolve())
    print(f"数据库已创建：{args.output.resolve()}")
    for name, count in counts.items():
        print(f"{name}: {count}")


if __name__ == "__main__":
    main()
