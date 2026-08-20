"""从业务数据 Excel 构建逐买家轮次的标注样本。"""

from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from emotion_agent.annotation_schemas import AnnotationCase, AnnotationMessage


REQUIRED_COLUMNS = {"会话ID", "消息序号", "角色", "message_text"}


def _normalize_text(value: Any) -> str:
    """将 Excel 单元格统一为适合 Prompt 的单行文本。"""
    return " ".join(str(value or "").split())


def build_annotation_cases(
    source_path: str | Path,
    sheet_name: str = "聊天记录",
) -> list[AnnotationCase]:
    """为每条买家消息生成从会话开头到当前时点的上下文。"""
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"工作簿中不存在工作表：{sheet_name}")
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            raise ValueError("聊天记录工作表为空")

        column_index = {str(name): index for index, name in enumerate(headers)}
        missing = REQUIRED_COLUMNS - column_index.keys()
        if missing:
            raise ValueError(f"聊天记录缺少必要列：{sorted(missing)}")

        conversations: dict[str, list[AnnotationMessage]] = defaultdict(list)
        for row in rows:
            conversation_id = _normalize_text(row[column_index["会话ID"]])
            role = _normalize_text(row[column_index["角色"]])
            text = _normalize_text(row[column_index["message_text"]])
            raw_seq = row[column_index["消息序号"]]
            if not conversation_id or not role or not text or raw_seq is None:
                continue
            conversations[conversation_id].append(
                AnnotationMessage(
                    seq=int(raw_seq),
                    role=role,
                    text=text,
                )
            )
    finally:
        workbook.close()

    cases: list[AnnotationCase] = []
    for conversation_id, conversation_messages in conversations.items():
        ordered_messages = sorted(conversation_messages, key=lambda item: item.seq)
        context: list[AnnotationMessage] = []
        buyer_turn = 0
        for message in ordered_messages:
            context.append(message)
            if message.role != "买家":
                continue
            buyer_turn += 1
            cases.append(
                AnnotationCase(
                    case_id=f"{conversation_id}_B{buyer_turn:02d}",
                    conversation_id=conversation_id,
                    buyer_turn=buyer_turn,
                    target_message_seq=message.seq,
                    target_text=message.text,
                    messages=[item.model_copy(deep=True) for item in context],
                )
            )
    return cases


def write_cases_jsonl(cases: list[AnnotationCase], output_path: str | Path) -> None:
    """将未标注样本写为 UTF-8 JSONL。"""
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as file:
        for case in cases:
            file.write(
                json.dumps(case.model_dump(), ensure_ascii=False, separators=(",", ":"))
                + "\n"
            )
